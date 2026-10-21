
import csv
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


MAX_PDF_PAGES = 5_000
MAX_PDF_TEXT_CHARACTERS = 5_000_000
MAX_DOCUMENT_TEXT_CHARACTERS = 5_000_000


def _check_text_limit(current: int, added: int) -> None:
    if current + added > MAX_DOCUMENT_TEXT_CHARACTERS:
        raise ValueError(
            "Extracted document text exceeds the 5 million character limit"
        )


def extract_pdf(file_path: str) -> list[dict]:
    pages = []
    total_characters = 0

    reader = PdfReader(file_path)

    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError("The PDF exceeds the 5,000 page processing limit")

    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()

        if not text:
            continue

        _check_text_limit(total_characters, len(text))
        total_characters += len(text)

        pages.append({"text": text, "page": page_number})

    return pages


def extract_docx(file_path: str) -> list[dict]:
    document = DocxDocument(file_path)
    paragraphs = []
    total_characters = 0

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        _check_text_limit(total_characters, len(text) + 1)
        total_characters += len(text) + 1
        paragraphs.append(text)

    text = "\n".join(paragraphs)

    return [{"text": text, "page": None}] if text else []


def extract_txt(file_path: str) -> list[dict]:
    path = Path(file_path)

    if path.stat().st_size > MAX_DOCUMENT_TEXT_CHARACTERS * 4:
        raise ValueError("The text file exceeds the processing size limit")

    text = path.read_text(encoding="utf-8", errors="ignore").strip()

    if len(text) > MAX_DOCUMENT_TEXT_CHARACTERS:
        raise ValueError("The text file exceeds the 5 million character limit")

    return [{"text": text, "page": None}] if text else []


def extract_csv(file_path: str) -> list[dict]:
    rows = []
    total_characters = 0

    with open(file_path, "r", encoding="utf-8-sig", errors="replace", newline="") as file:
        reader = csv.reader(file)

        for row in reader:
            line = " | ".join(row)
            _check_text_limit(total_characters, len(line) + 1)
            total_characters += len(line) + 1

            if line.strip():
                rows.append(line)

    text = "\n".join(rows).strip()
    return [{"text": text, "page": None}] if text else []


def extract_excel(file_path: str) -> list[dict]:
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    sheet_texts = []
    total_characters = 0

    try:
        for sheet in workbook.worksheets:
            lines = []

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value) if value is not None else ""
                    for value in row
                ]
                line = " | ".join(values).rstrip()

                if not line.strip():
                    continue

                _check_text_limit(total_characters, len(line) + 1)
                total_characters += len(line) + 1
                lines.append(line)

            if lines:
                sheet_texts.append(
                    f"Sheet: {sheet.title}\n" + "\n".join(lines)
                )
    finally:
        workbook.close()

    text = "\n\n".join(sheet_texts).strip()
    return [{"text": text, "page": None}] if text else []


def extract_pptx(file_path: str) -> list[dict]:
    presentation = Presentation(file_path)
    slides = []
    total_characters = 0

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_texts = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            text = shape.text.strip()

            if text:
                _check_text_limit(total_characters, len(text) + 1)
                total_characters += len(text) + 1
                slide_texts.append(text)

        text = "\n".join(slide_texts).strip()

        if text:
            slides.append({"text": text, "page": slide_number})

    return slides
