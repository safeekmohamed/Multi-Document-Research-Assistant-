from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    UploadFile,
)
from fastapi.routing import APIRoute
from fastapi.responses import StreamingResponse
from starlette.background import BackgroundTask
from starlette.exceptions import HTTPException as StarletteHTTPException
from starlette.responses import JSONResponse

from bson import ObjectId
from urllib.parse import quote, urlsplit
import logging
import threading
import time
from io import BytesIO
import tempfile
from pathlib import Path
import requests
import cloudinary.utils
from openpyxl import load_workbook
from starlette.concurrency import run_in_threadpool

from app.auth.jwt import get_current_user
from app.database.mongodb import (
    documents_collection,
)
from app.documents.document_service import (
    delete_user_document,
    list_user_documents,
    process_uploaded_document,
)


MAX_VIEW_SIZE = 50 * 1024 * 1024
MAX_UPLOAD_FILES = 10
MAX_UPLOAD_FILE_SIZE = 20 * 1024 * 1024
MAX_UPLOAD_REQUEST_SIZE = 50 * 1024 * 1024
UPLOAD_COPY_CHUNK_SIZE = 1024 * 1024
logger = logging.getLogger(__name__)
_UPLOAD_PROCESS_LOCK = threading.Lock()


def _process_document_serially(**kwargs):
    with _UPLOAD_PROCESS_LOCK:
        return process_uploaded_document(**kwargs)


class BoundedUploadRoute(APIRoute):
    def get_route_handler(self):
        route_handler = super().get_route_handler()

        async def handle_request(request):
            if self.name == "upload_documents" and request.method == "POST":
                content_length = request.headers.get("content-length", "")
                if content_length.isdigit() and int(content_length) > MAX_UPLOAD_REQUEST_SIZE + 1024 * 1024:
                    return JSONResponse(
                        status_code=413,
                        content={"detail": "The combined upload must be 50 MB or smaller."},
                    )
                try:
                    await request.form(
                        max_files=MAX_UPLOAD_FILES,
                        max_fields=MAX_UPLOAD_FILES,
                    )
                except StarletteHTTPException as error:
                    if "too many files" not in str(error.detail).lower():
                        raise
                    return JSONResponse(
                        status_code=413,
                        content={"detail": f"Upload no more than {MAX_UPLOAD_FILES} files at once."},
                    )
            return await route_handler(request)

        return handle_request


router = APIRouter(
    prefix="/documents",
    tags=["Documents"],
    route_class=BoundedUploadRoute,
)


def _get_upload_size(file: UploadFile) -> int:
    if file.size is not None:
        return int(file.size)
    current_position = file.file.tell()
    file.file.seek(0, 2)
    size = file.file.tell()
    file.file.seek(current_position)
    return size


async def _stage_upload(file: UploadFile) -> tuple[Path, int]:
    suffix = Path(file.filename or "document").suffix.lower()
    staged_file = tempfile.NamedTemporaryFile(
        mode="wb",
        suffix=suffix,
        delete=False,
    )
    staged_path = Path(staged_file.name)
    size = 0
    started_at = time.perf_counter()
    try:
        await file.seek(0)
        while True:
            block = await file.read(UPLOAD_COPY_CHUNK_SIZE)
            if not block:
                break
            size += len(block)
            if size > MAX_UPLOAD_FILE_SIZE:
                raise HTTPException(
                    status_code=413,
                    detail="Each uploaded file must be 20 MB or smaller.",
                )
            staged_file.write(block)
        if size == 0:
            raise HTTPException(
                status_code=400,
                detail="Uploaded file is empty",
            )
        staged_file.close()
        logger.info(
            "Upload staging completed: bytes=%d elapsed_seconds=%.3f",
            size,
            time.perf_counter() - started_at,
        )
        return staged_path, size
    except Exception:
        staged_file.close()
        staged_path.unlink(missing_ok=True)
        raise


def _owned_file_document(document_id: str, current_user: dict):
    document_filter = {"user_id": str(current_user["_id"])}
    if ObjectId.is_valid(document_id):
        document_filter["_id"] = ObjectId(document_id)
    else:
        document_filter["source_id"] = document_id
    document = documents_collection.find_one(document_filter)
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    if document.get("source_type", "file") != "file":
        raise HTTPException(status_code=400, detail="This source is not an uploaded file")
    return document


def _cloudinary_response(document):
    public_id = document.get("cloudinary_public_id")
    cloudinary_url = document.get("cloudinary_url")
    if not public_id and cloudinary_url:
        parsed_url = urlsplit(cloudinary_url)
        if parsed_url.scheme == "https" and parsed_url.hostname == "res.cloudinary.com":
            path_parts = parsed_url.path.strip("/").split("/")
            try:
                upload_index = path_parts.index("upload")
                if path_parts[upload_index - 1] == "raw":
                    resource_parts = path_parts[upload_index + 1 :]
                    if resource_parts and resource_parts[0].startswith("v") and resource_parts[0][1:].isdigit():
                        resource_parts = resource_parts[1:]
                    public_id = "/".join(resource_parts)
            except (ValueError, IndexError):
                public_id = None
    if not public_id:
        raise HTTPException(status_code=404, detail="Cloudinary file reference not found")

    file_response = None
    try:
        download_url = cloudinary.utils.private_download_url(
            public_id,
            str(document.get("file_type") or "pdf").lstrip("."),
            resource_type="raw",
            type="upload",
            expires_at=int(time.time()) + 300,
            secure=True,
        )
        parsed_download_url = urlsplit(download_url)
        if parsed_download_url.scheme != "https" or parsed_download_url.hostname != "api.cloudinary.com":
            raise ValueError("Cloudinary generated an invalid download URL")
        file_response = requests.get(download_url, stream=True, timeout=30, allow_redirects=False)
        if file_response.status_code == 404:
            file_response.close()
            raise HTTPException(status_code=404, detail="Stored Cloudinary file was not found")
        if not 200 <= file_response.status_code < 300:
            raise requests.HTTPError("Cloudinary returned a non-success response", response=file_response)
        content_length = file_response.headers.get("Content-Length")
        if content_length and content_length.isdigit() and int(content_length) > MAX_VIEW_SIZE:
            file_response.close()
            raise HTTPException(status_code=413, detail="This file exceeds the 50 MB preview limit.")
        return file_response
    except HTTPException:
        raise
    except Exception as error:
        response = file_response or getattr(error, "response", None)
        if response is not None:
            response.close()
        raise HTTPException(status_code=502, detail="Unable to retrieve the stored document") from error


@router.post("/upload")
async def upload_documents(
    files: list[UploadFile] = File(
        description=(
            "Upload one or more PDF, DOCX, or TXT files"
        )
    ),
    current_user: dict = Depends(
        get_current_user
    ),
):
    if not files:
        raise HTTPException(
            status_code=400,
            detail="No files were uploaded",
        )

    staged_paths = []
    processing_stage = "validation"
    try:
        if len(files) > MAX_UPLOAD_FILES:
            raise HTTPException(
                status_code=413,
                detail=f"Upload no more than {MAX_UPLOAD_FILES} files at once.",
            )

        declared_sizes = [_get_upload_size(file) for file in files]
        if any(size > MAX_UPLOAD_FILE_SIZE for size in declared_sizes):
            raise HTTPException(
                status_code=413,
                detail="Each uploaded file must be 20 MB or smaller.",
            )
        if sum(declared_sizes) > MAX_UPLOAD_REQUEST_SIZE:
            raise HTTPException(
                status_code=413,
                detail="The combined upload must be 50 MB or smaller.",
            )

        logger.info(
            "Document upload accepted: file_count=%d total_bytes=%d",
            len(files),
            sum(declared_sizes),
        )
        results = []
        actual_total_size = 0
        for file in files:
            filename = file.filename or "document"
            processing_stage = "staging"
            staged_path, actual_size = await _stage_upload(file)
            staged_paths.append(staged_path)
            await file.close()
            actual_total_size += actual_size
            if actual_total_size > MAX_UPLOAD_REQUEST_SIZE:
                raise HTTPException(
                    status_code=413,
                    detail="The combined upload must be 50 MB or smaller.",
                )

            try:
                processing_stage = "document_processing"
                processing_started_at = time.perf_counter()
                result = await run_in_threadpool(
                    _process_document_serially,
                    user_id=str(current_user["_id"]),
                    filename=filename,
                    file_path=staged_path,
                )
                results.append(result)
                logger.info(
                    "Document upload completed: bytes=%d chunks=%d elapsed_seconds=%.3f",
                    actual_size,
                    int(result.get("chunk_count", 0)),
                    time.perf_counter() - processing_started_at,
                )
            except ValueError as error:
                raise HTTPException(
                    status_code=400,
                    detail=str(error),
                ) from error
            except Exception as error:
                logger.error(
                    "Document upload failed: stage=%s error_type=%s",
                    processing_stage,
                    type(error).__name__,
                )
                raise HTTPException(
                    status_code=500,
                    detail="Document processing failed. Try a smaller supported file or retry later.",
                ) from error
            finally:
                staged_path.unlink(missing_ok=True)
                staged_paths.remove(staged_path)

        return {
            "message": "Documents processed successfully",
            "documents": results,
        }
    finally:
        for staged_path in staged_paths:
            staged_path.unlink(missing_ok=True)
        for file in files:
            await file.close()


@router.get("/")
def get_documents(
    current_user: dict = Depends(
        get_current_user
    ),
):
    return {
        "documents": list_user_documents(
            str(
                current_user["_id"]
            )
        )
    }


@router.get("/view/{document_id}")
def view_document(
    document_id: str,
    current_user: dict = Depends(
        get_current_user
    ),
):
    document = _owned_file_document(document_id, current_user)
    file_response = _cloudinary_response(document)
    spool = tempfile.SpooledTemporaryFile(max_size=2 * 1024 * 1024, mode="w+b")
    total_size = 0
    try:
        for chunk in file_response.iter_content(chunk_size=64 * 1024):
            if not chunk:
                continue
            total_size += len(chunk)
            if total_size > MAX_VIEW_SIZE:
                raise HTTPException(status_code=413, detail="This file exceeds the 50 MB preview limit.")
            spool.write(chunk)
    except HTTPException:
        spool.close()
        file_response.close()
        raise
    except Exception as error:
        spool.close()
        file_response.close()
        raise HTTPException(status_code=502, detail="Unable to retrieve the stored document") from error
    finally:
        file_response.close()
    spool.seek(0)

    filename = quote(
        str(document.get("filename") or "document"),
        safe="",
    )
    headers = {
        "Content-Disposition": f"inline; filename*=UTF-8''{filename}",
        "X-Content-Type-Options": "nosniff",
        "X-Document-File-Type": str(document.get("file_type") or ""),
        "X-Document-Filename": quote(str(document.get("filename") or "document"), safe=""),
    }
    headers["Content-Length"] = str(total_size)

    return StreamingResponse(
        spool,
        media_type=file_response.headers.get(
            "Content-Type",
            "application/octet-stream",
        ),
        headers=headers,
        background=BackgroundTask(spool.close),
    )


@router.get("/preview/{document_id}")
def preview_document_data(
    document_id: str,
    current_user: dict = Depends(get_current_user),
):
    document = _owned_file_document(document_id, current_user)
    if str(document.get("file_type") or "").lower() != "xlsx":
        raise HTTPException(status_code=415, detail="Structured preview is available for XLSX files only")
    file_response = _cloudinary_response(document)
    try:
        content_buffer = bytearray()
        for chunk in file_response.iter_content(chunk_size=64 * 1024):
            content_buffer.extend(chunk)
            if len(content_buffer) > MAX_VIEW_SIZE:
                raise HTTPException(status_code=413, detail="This file exceeds the 50 MB preview limit.")
        content = bytes(content_buffer)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=502, detail="Unable to retrieve the stored document") from error
    finally:
        file_response.close()
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheets = []
        was_truncated = False
        for sheet in workbook.worksheets:
            sheet_rows = []
            max_row = min(sheet.max_row or 0, 501)
            max_column = min(sheet.max_column or 0, 51)
            was_truncated = was_truncated or (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 50
            for row in sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_column, values_only=True):
                values = ["" if value is None else str(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                sheet_rows.append(values)
            sheets.append({"name": sheet.title, "rows": sheet_rows})
        workbook.close()
    except Exception as error:
        raise HTTPException(status_code=422, detail="The stored XLSX file could not be parsed") from error
    return {
        "filename": document.get("filename") or "document.xlsx",
        "sheets": sheets,
        "truncated": was_truncated,
    }


@router.delete("/{document_id}")
def delete_document(
    document_id: str,
    current_user: dict = Depends(
        get_current_user
    ),
):
    try:
        document = delete_user_document(
            user_id=str(
                current_user["_id"]
            ),
            document_id=document_id,
        )

        return {
            "message": (
                "Document deleted successfully"
            ),
            "document": {
                "id": str(
                    document["_id"]
                ),
                "filename": document.get(
                    "filename",
                    "",
                ),
            },
        }

    except ValueError as error:
        raise HTTPException(
            status_code=400,
            detail=str(error),
        )

    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=(
                f"Document deletion failed: "
                f"{error}"
            ),
        )
